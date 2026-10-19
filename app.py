import os
import sys
import uuid
import threading
import webbrowser
import tempfile
from pathlib import Path

from flask import Flask, render_template, request, send_file, jsonify
from PIL import Image
from werkzeug.utils import secure_filename


def resource_path(relative_path):
    if getattr(sys, "frozen", False):
        return Path(sys._MEIPASS) / relative_path
    return Path(__file__).parent / relative_path


app = Flask(
    __name__,
    template_folder=str(resource_path("templates")),
)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # 50 MB

DATA_DIR = Path(tempfile.gettempdir()) / "fileflow_data"
UPLOAD_DIR = DATA_DIR / "uploads"
OUTPUT_DIR = DATA_DIR / "outputs"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "bmp", "webp", "gif", "tiff", "tif", "ico"}

CONVERSION_MAP = {
    "png": IMAGE_EXTENSIONS - {"png"},
    "jpg": IMAGE_EXTENSIONS - {"jpg"},
    "jpeg": IMAGE_EXTENSIONS - {"jpeg"},
    "bmp": IMAGE_EXTENSIONS - {"bmp"},
    "webp": IMAGE_EXTENSIONS - {"webp"},
    "gif": IMAGE_EXTENSIONS - {"gif"},
    "tiff": IMAGE_EXTENSIONS - {"tiff"},
    "tif": IMAGE_EXTENSIONS - {"tif"},
    "ico": IMAGE_EXTENSIONS - {"ico"},
    "pdf": {"docx"},
    "docx": {"pdf"},
    "xlsx": {"csv"},
    "csv": {"xlsx"},
}


def get_targets(ext: str):
    ext = ext.lower().lstrip(".")
    return sorted(CONVERSION_MAP.get(ext, []))


def convert_image(src: Path, target_ext: str) -> Path:
    out = OUTPUT_DIR / f"{src.stem}.{target_ext}"
    img = Image.open(src)
    if img.mode == "RGBA" and target_ext in ("jpg", "jpeg", "bmp", "ico"):
        bg = Image.new("RGB", img.size, (255, 255, 255))
        bg.paste(img, mask=img.split()[3])
        img = bg
    elif img.mode != "RGB" and target_ext in ("jpg", "jpeg", "bmp"):
        img = img.convert("RGB")
    img.save(out)
    return out


def convert_pdf_to_docx(src: Path) -> Path:
    from pdf2docx import Converter
    out = OUTPUT_DIR / f"{src.stem}.docx"
    cv = Converter(str(src))
    cv.convert(str(out))
    cv.close()
    return out


def convert_docx_to_pdf(src: Path) -> Path:
    from docx import Document
    from fpdf import FPDF

    doc = Document(str(src))
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for para in doc.paragraphs:
        pdf.multi_cell(0, 7, para.text)
        pdf.ln(2)
    out = OUTPUT_DIR / f"{src.stem}.pdf"
    pdf.output(str(out))
    return out


def convert_xlsx_to_csv(src: Path) -> Path:
    import openpyxl
    import csv
    wb = openpyxl.load_workbook(str(src))
    ws = wb.active
    out = OUTPUT_DIR / f"{src.stem}.csv"
    with open(out, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(row)
    return out


def convert_csv_to_xlsx(src: Path) -> Path:
    import openpyxl
    import csv
    wb = openpyxl.Workbook()
    ws = wb.active
    with open(src, "r", encoding="utf-8") as f:
        reader = csv.reader(f)
        for row in reader:
            ws.append(row)
    out = OUTPUT_DIR / f"{src.stem}.xlsx"
    wb.save(str(out))
    return out


def remove_bg(src: Path) -> Path:
    from rembg import remove
    inp = Image.open(src)
    result = remove(inp)
    out = OUTPUT_DIR / f"{src.stem}_no_bg.png"
    result.save(out)
    return out


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/targets", methods=["POST"])
def api_targets():
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "Nenhum arquivo enviado"}), 400
    ext = Path(f.filename).suffix.lower().lstrip(".")
    targets = get_targets(ext)
    return jsonify({"source_ext": ext, "targets": targets})


@app.route("/api/convert", methods=["POST"])
def api_convert():
    f = request.files.get("file")
    target = request.form.get("target", "").lower()
    if not f:
        return jsonify({"error": "Nenhum arquivo enviado"}), 400

    uid = uuid.uuid4().hex[:8]
    filename = secure_filename(f.filename)
    src = UPLOAD_DIR / f"{uid}_{filename}"
    f.save(src)

    ext = Path(filename).suffix.lower().lstrip(".")

    try:
        if ext in IMAGE_EXTENSIONS and target in IMAGE_EXTENSIONS:
            out = convert_image(src, target)
        elif ext == "pdf" and target == "docx":
            out = convert_pdf_to_docx(src)
        elif ext == "docx" and target == "pdf":
            out = convert_docx_to_pdf(src)
        elif ext == "xlsx" and target == "csv":
            out = convert_xlsx_to_csv(src)
        elif ext == "csv" and target == "xlsx":
            out = convert_csv_to_xlsx(src)
        else:
            return jsonify({"error": f"Conversão {ext} → {target} não suportada"}), 400

        return send_file(out, as_attachment=True, download_name=out.name)
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        src.unlink(missing_ok=True)


@app.route("/api/remove-bg", methods=["POST"])
def api_remove_bg():
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "Nenhum arquivo enviado"}), 400

    uid = uuid.uuid4().hex[:8]
    filename = secure_filename(f.filename)
    src = UPLOAD_DIR / f"{uid}_{filename}"
    f.save(src)

    try:
        out = remove_bg(src)
        return send_file(out, as_attachment=True, download_name=out.name)
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        src.unlink(missing_ok=True)


def open_browser():
    webbrowser.open("http://localhost:5000")


if __name__ == "__main__":
    if getattr(sys, "frozen", False):
        threading.Timer(1.5, open_browser).start()
        app.run(host="127.0.0.1", port=5000, debug=False)
    else:
        app.run(debug=True, port=5000)
